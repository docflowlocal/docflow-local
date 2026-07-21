#!/usr/bin/env python3
"""DocFlow Local — local-first batch document generation MVP."""

from __future__ import annotations

import ast
import base64
import csv
import hashlib
import io
import json
import math
import os
import re
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Tuple

# The preinstalled Pillow build expects a newer NumPy typing surface. This
# compatibility shim is harmless on current NumPy and keeps the MVP runnable
# in older managed Python environments.
try:
    import numpy.typing as _npt

    if not hasattr(_npt, "NDArray"):
        class _NDArray:
            def __class_getitem__(cls, item):
                return object

        _npt.NDArray = _NDArray
except Exception:
    pass

# ReportLab 4 passes the Python 3.9+ ``usedforsecurity`` flag to hashlib.
# Ignore it on Python 3.8 so one codebase works on both runtime generations.
_original_md5 = hashlib.md5
def _compatible_md5(*args, **kwargs):
    kwargs.pop("usedforsecurity", None)
    return _original_md5(*args, **kwargs)
hashlib.md5 = _compatible_md5

from flask import Flask, jsonify, request, send_file
from openpyxl import load_workbook
from reportlab.graphics import renderPDF
from reportlab.graphics.barcode import qr
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, Table, TableStyle


ROOT = Path(__file__).resolve().parent
app = Flask(__name__, static_folder=str(ROOT / "static"), static_url_path="")
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024


def _register_fonts() -> Tuple[str, str]:
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/Library/Fonts/Arial Unicode.ttf",
        "C:/Windows/Fonts/msyh.ttc",
        "C:/Windows/Fonts/simhei.ttf",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for candidate in candidates:
        if os.path.exists(candidate):
            try:
                pdfmetrics.registerFont(TTFont("DocFlowCJK", candidate, subfontIndex=0))
                return "DocFlowCJK", "DocFlowCJK"
            except Exception:
                continue
    return "Helvetica", "Helvetica-Bold"


FONT_REGULAR, FONT_BOLD = _register_fonts()


def _decode_csv(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030", "utf-16"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _clean_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def parse_tabular(filename: str, raw: bytes) -> Tuple[List[str], List[Dict[str, Any]]]:
    suffix = Path(filename).suffix.lower()
    matrix: List[List[Any]] = []
    if suffix == ".csv":
        matrix = list(csv.reader(io.StringIO(_decode_csv(raw))))
    elif suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.active
        matrix = [[_clean_value(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
    else:
        raise ValueError("仅支持 CSV、XLSX 或 XLSM 数据文件")

    matrix = [row for row in matrix if any(str(cell).strip() for cell in row)]
    if not matrix:
        return [], []
    headers = [str(value).strip() or f"字段_{index + 1}" for index, value in enumerate(matrix[0])]
    rows = []
    for values in matrix[1:]:
        padded = list(values) + [""] * (len(headers) - len(values))
        rows.append({header: _clean_value(value) for header, value in zip(headers, padded[: len(headers)])})
    return headers, rows


def extract_docx_fields(raw: bytes) -> List[str]:
    fields = set()
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        for name in archive.namelist():
            if not name.startswith("word/") or not name.endswith(".xml"):
                continue
            xml = archive.read(name).decode("utf-8", errors="ignore")
            # Word may split visible text across multiple XML nodes.
            visible = re.sub(r"<[^>]+>", "", xml)
            for match in re.findall(r"\{\{\s*([^{}]+?)\s*\}\}", visible):
                fields.add(match.strip())
    return sorted(fields)


ALLOWED_NODES = (
    ast.Expression, ast.BinOp, ast.UnaryOp, ast.Constant, ast.Name, ast.Load,
    ast.Add, ast.Sub, ast.Mult, ast.Div, ast.Mod, ast.Pow, ast.USub, ast.UAdd,
    ast.Call,
)
ALLOWED_FUNCTIONS = {"round": round, "min": min, "max": max, "abs": abs}


def _number(value: Any) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    normalized = str(value).replace(",", "").replace("¥", "").replace("￥", "").strip()
    if normalized.endswith("%"):
        return float(normalized[:-1] or 0) / 100
    return float(normalized or 0)


def evaluate_formula(expression: str, row: Dict[str, Any]) -> float:
    names = {key: _number(value) for key, value in row.items() if str(value).strip()}
    tree = ast.parse(expression, mode="eval")
    for node in ast.walk(tree):
        if not isinstance(node, ALLOWED_NODES):
            raise ValueError("公式包含不支持的运算")
        if isinstance(node, ast.Call) and not (
            isinstance(node.func, ast.Name) and node.func.id in ALLOWED_FUNCTIONS
        ):
            raise ValueError("公式函数不受支持")
    return float(eval(compile(tree, "<formula>", "eval"), {"__builtins__": {}}, {**ALLOWED_FUNCTIONS, **names}))


def apply_rules(rows: Iterable[Dict[str, Any]], computed_fields: Iterable[Dict[str, str]]) -> List[Dict[str, Any]]:
    output = []
    for source in rows:
        row = dict(source)
        for field in computed_fields or []:
            name = field.get("name", "").strip()
            expression = field.get("expression", "").strip()
            if name and expression:
                try:
                    row[name] = round(evaluate_formula(expression, row), 2)
                except Exception:
                    row[name] = ""
        output.append(row)
    return output


def validate_rows(rows: List[Dict[str, Any]], required_fields: Iterable[str]) -> Dict[str, Any]:
    required = [field for field in required_fields if field]
    issues = []
    valid_indexes = []
    for index, row in enumerate(rows):
        missing = [field for field in required if str(row.get(field, "")).strip() == ""]
        if missing:
            issues.append({
                "row": index + 2,
                "record": str(row.get("客户简称") or row.get("客户名称") or f"第 {index + 1} 条"),
                "missing": missing,
            })
        else:
            valid_indexes.append(index)
    return {
        "total": len(rows),
        "valid": len(valid_indexes),
        "invalid": len(issues),
        "issues": issues,
        "validIndexes": valid_indexes,
    }


def safe_component(value: Any, fallback: str = "未命名") -> str:
    text = re.sub(r"[\\/:*?\"<>|\x00-\x1f]", "-", str(value or "").strip())
    text = re.sub(r"\s+", " ", text).strip(" .")
    return (text or fallback)[:96]


def render_pattern(pattern: str, row: Dict[str, Any]) -> str:
    return re.sub(
        r"\{\{\s*([^{}]+?)\s*\}\}",
        lambda match: safe_component(row.get(match.group(1).strip(), match.group(1).strip())),
        pattern,
    )


def _money(value: Any) -> str:
    try:
        return f"¥{_number(value):,.2f}"
    except Exception:
        return f"¥{value or 0}"


def _draw_qr(pdf: canvas.Canvas, value: str, x: float, y: float, size: float = 22 * mm) -> None:
    widget = qr.QrCodeWidget(value)
    bounds = widget.getBounds()
    width, height = bounds[2] - bounds[0], bounds[3] - bounds[1]
    drawing = Drawing(size, size, transform=[size / width, 0, 0, size / height, 0, 0])
    drawing.add(widget)
    renderPDF.draw(drawing, pdf, x, y)


def _paragraph(text: Any, size: int = 9, color=colors.HexColor("#324152"), align=TA_LEFT):
    return Paragraph(
        str(text or "—"),
        ParagraphStyle(
            "df", fontName=FONT_REGULAR, fontSize=size, leading=size * 1.45,
            textColor=color, alignment=align,
        ),
    )


def _image_from_data_url(data_url: str):
    if not data_url or "," not in data_url:
        return None
    try:
        _, encoded = data_url.split(",", 1)
        return ImageReader(io.BytesIO(base64.b64decode(encoded)))
    except Exception:
        return None


def make_quote_pdf(row: Dict[str, Any], output: io.BytesIO, signature_data: str = "") -> None:
    pdf = canvas.Canvas(output, pagesize=A4)
    width, height = A4
    navy = colors.HexColor("#12243A")
    teal = colors.HexColor("#0C8B86")
    muted = colors.HexColor("#697789")
    pale = colors.HexColor("#EDF7F6")

    pdf.setFillColor(navy)
    pdf.rect(0, height - 38 * mm, width, 38 * mm, fill=1, stroke=0)
    pdf.setFillColor(teal)
    pdf.roundRect(18 * mm, height - 27 * mm, 16 * mm, 16 * mm, 3 * mm, fill=1, stroke=0)
    pdf.setFillColor(colors.white)
    pdf.setFont("Helvetica-Bold", 13)
    pdf.drawCentredString(26 * mm, height - 21.5 * mm, "DF")
    pdf.setFont(FONT_BOLD, 20)
    pdf.drawString(41 * mm, height - 17 * mm, "报价单")
    pdf.setFont(FONT_REGULAR, 8)
    pdf.setFillColor(colors.HexColor("#B8C5D3"))
    pdf.drawString(41 * mm, height - 23 * mm, "DOCFLOW LOCAL · QUOTATION")
    pdf.setFillColor(colors.white)
    pdf.setFont(FONT_REGULAR, 9)
    pdf.drawRightString(width - 18 * mm, height - 17 * mm, f"编号  {row.get('报价编号', '—')}")
    pdf.drawRightString(width - 18 * mm, height - 23 * mm, f"日期  {row.get('报价日期', datetime.now().strftime('%Y-%m-%d'))}")

    y = height - 52 * mm
    pdf.setFillColor(muted)
    pdf.setFont(FONT_REGULAR, 8)
    pdf.drawString(18 * mm, y, "报价对象")
    pdf.drawString(112 * mm, y, "联系信息")
    pdf.setFillColor(navy)
    pdf.setFont(FONT_BOLD, 13)
    pdf.drawString(18 * mm, y - 7 * mm, str(row.get("客户名称") or row.get("客户简称") or "—"))
    pdf.setFont(FONT_REGULAR, 9)
    pdf.drawString(112 * mm, y - 6 * mm, f"联系人：{row.get('联系人', '—')}")
    pdf.drawString(112 * mm, y - 12 * mm, f"邮箱：{row.get('邮箱', '—')}")

    quantity = _number(row.get("数量", 0))
    unit_price = _number(row.get("单价", 0))
    subtotal = _number(row.get("小计", quantity * unit_price))
    discount = _number(row.get("优惠", 0))
    tax_rate = _number(row.get("税率", "13%"))
    tax = _number(row.get("税额", max(subtotal - discount, 0) * tax_rate))
    total = _number(row.get("含税总额", max(subtotal - discount, 0) + tax))
    table_data = [
        [_paragraph("项目", 8, colors.white), _paragraph("规格/说明", 8, colors.white), _paragraph("数量", 8, colors.white, TA_RIGHT), _paragraph("单价", 8, colors.white, TA_RIGHT), _paragraph("金额", 8, colors.white, TA_RIGHT)],
        [_paragraph(row.get("产品名称", "专业服务")), _paragraph(row.get("产品说明", "按双方确认范围交付")), _paragraph(str(row.get("数量", 1)), align=TA_RIGHT), _paragraph(_money(unit_price), align=TA_RIGHT), _paragraph(_money(subtotal), align=TA_RIGHT)],
    ]
    table = Table(table_data, colWidths=[45 * mm, 68 * mm, 18 * mm, 27 * mm, 28 * mm], rowHeights=[10 * mm, 20 * mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), navy),
        ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#F7F9FB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#DDE3E9")),
    ]))
    table.wrapOn(pdf, width, height)
    table.drawOn(pdf, 12 * mm, y - 53 * mm)

    summary_y = y - 70 * mm
    labels = [("小计", subtotal)]
    if discount > 0:
        labels.append(("优惠", -discount))
    labels.append((f"税额（{tax_rate:.0%}）", tax))
    pdf.setFont(FONT_REGULAR, 9)
    for label, value in labels:
        pdf.setFillColor(muted)
        pdf.drawRightString(width - 51 * mm, summary_y, label)
        pdf.setFillColor(navy)
        pdf.drawRightString(width - 18 * mm, summary_y, _money(value))
        summary_y -= 7 * mm
    pdf.setStrokeColor(colors.HexColor("#D8E0E7"))
    pdf.line(width - 93 * mm, summary_y + 3 * mm, width - 18 * mm, summary_y + 3 * mm)
    pdf.setFillColor(teal)
    pdf.setFont(FONT_BOLD, 13)
    pdf.drawRightString(width - 51 * mm, summary_y - 3 * mm, "含税总额")
    pdf.drawRightString(width - 18 * mm, summary_y - 3 * mm, _money(total))

    note_y = 77 * mm
    pdf.setFillColor(pale)
    pdf.roundRect(18 * mm, note_y, width - 36 * mm, 36 * mm, 3 * mm, fill=1, stroke=0)
    pdf.setFillColor(navy)
    pdf.setFont(FONT_BOLD, 9)
    pdf.drawString(24 * mm, note_y + 26 * mm, "条款与说明")
    pdf.setFont(FONT_REGULAR, 8.5)
    pdf.setFillColor(muted)
    notes = str(row.get("备注") or "本报价有效期 30 天；交付范围与付款方式以双方确认内容为准。")
    pdf.drawString(24 * mm, note_y + 17 * mm, notes[:58])
    pdf.drawString(24 * mm, note_y + 10 * mm, "本文件由 DocFlow Local 在本机生成，数据未上传至云端。")

    qr_value = str(row.get("二维码内容") or f"quote:{row.get('报价编号', '')}|customer:{row.get('客户简称', '')}|total:{total:.2f}")
    _draw_qr(pdf, qr_value, width - 43 * mm, 30 * mm)
    pdf.setFillColor(muted)
    pdf.setFont(FONT_REGULAR, 7)
    pdf.drawCentredString(width - 32 * mm, 26 * mm, "扫码核验报价信息")
    pdf.setStrokeColor(colors.HexColor("#9BA8B6"))
    pdf.line(18 * mm, 43 * mm, 78 * mm, 43 * mm)
    pdf.line(91 * mm, 43 * mm, 151 * mm, 43 * mm)
    signature = _image_from_data_url(signature_data)
    if signature:
        try:
            pdf.drawImage(signature, 106 * mm, 45 * mm, 30 * mm, 14 * mm, preserveAspectRatio=True, anchor="c", mask="auto")
        except Exception:
            pass
    pdf.setFillColor(muted)
    pdf.drawString(18 * mm, 38 * mm, "客户确认 / 日期")
    pdf.drawString(91 * mm, 38 * mm, "授权签名 / 日期")
    pdf.setFillColor(colors.HexColor("#9AA6B2"))
    pdf.setFont(FONT_REGULAR, 7)
    pdf.drawString(18 * mm, 16 * mm, "DOCFLOW LOCAL · PRIVATE BY DESIGN")
    pdf.drawRightString(width - 18 * mm, 16 * mm, "第 1 页 / 共 1 页")
    pdf.save()


def make_attachment_pdf(row: Dict[str, Any], output: io.BytesIO) -> None:
    pdf = canvas.Canvas(output, pagesize=A4)
    width, height = A4
    navy = colors.HexColor("#12243A")
    teal = colors.HexColor("#0C8B86")
    pdf.setFillColor(navy)
    pdf.rect(0, height - 26 * mm, width, 26 * mm, fill=1, stroke=0)
    pdf.setFillColor(colors.white)
    pdf.setFont(FONT_BOLD, 15)
    pdf.drawString(18 * mm, height - 17 * mm, "交付附件 · 项目明细")
    pdf.setFont(FONT_REGULAR, 8)
    pdf.drawRightString(width - 18 * mm, height - 16 * mm, str(row.get("报价编号", "—")))
    pdf.setFillColor(teal)
    pdf.rect(18 * mm, height - 40 * mm, 2.2 * mm, 8 * mm, fill=1, stroke=0)
    pdf.setFillColor(navy)
    pdf.setFont(FONT_BOLD, 11)
    pdf.drawString(25 * mm, height - 38 * mm, str(row.get("产品名称", "项目范围")))
    pdf.setFont(FONT_REGULAR, 9)
    pdf.setFillColor(colors.HexColor("#5B6878"))
    lines = [
        f"客户：{row.get('客户名称') or row.get('客户简称') or '—'}",
        f"项目说明：{row.get('产品说明') or '按确认的需求清单执行。'}",
        f"交付周期：{row.get('交付周期') or '合同确认后 10 个工作日'}",
        f"负责人：{row.get('负责人') or '项目交付组'}",
        f"备注：{row.get('备注') or '具体里程碑以项目启动会确认为准。'}",
    ]
    y = height - 55 * mm
    for line in lines:
        pdf.drawString(25 * mm, y, line[:76])
        y -= 10 * mm
    pdf.setStrokeColor(colors.HexColor("#DDE3E9"))
    pdf.roundRect(18 * mm, 35 * mm, width - 36 * mm, 42 * mm, 3 * mm, fill=0, stroke=1)
    pdf.setFillColor(navy)
    pdf.setFont(FONT_BOLD, 9)
    pdf.drawString(25 * mm, 66 * mm, "交付检查")
    pdf.setFont(FONT_REGULAR, 8.5)
    for index, text in enumerate(["资料字段完整", "主文件可正常打开", "附件与客户编号一致", "交付包目录符合命名规则"]):
        yy = 57 * mm - index * 7 * mm
        pdf.rect(25 * mm, yy - 1.5 * mm, 3 * mm, 3 * mm, fill=0, stroke=1)
        pdf.drawString(32 * mm, yy, text)
    pdf.setFont(FONT_REGULAR, 7)
    pdf.setFillColor(colors.HexColor("#9AA6B2"))
    pdf.drawString(18 * mm, 16 * mm, "由 DocFlow Local 在本机生成")
    pdf.drawRightString(width - 18 * mm, 16 * mm, "附件 1")
    pdf.save()


@app.route("/", methods=["GET"])
def index():
    return app.send_static_file("index.html")


@app.route("/api/health", methods=["GET"])
def health():
    return jsonify({"ok": True, "mode": "local", "version": "0.1.0"})


@app.route("/api/import", methods=["POST"])
def import_data():
    upload = request.files.get("file")
    if not upload:
        return jsonify({"error": "请选择数据文件"}), 400
    try:
        headers, rows = parse_tabular(upload.filename or "data.csv", upload.read())
        return jsonify({"filename": upload.filename, "headers": headers, "rows": rows, "count": len(rows)})
    except Exception as exc:
        return jsonify({"error": f"读取失败：{exc}"}), 400


@app.route("/api/template", methods=["POST"])
def inspect_template():
    upload = request.files.get("file")
    if not upload:
        return jsonify({"error": "请选择模板文件"}), 400
    suffix = Path(upload.filename or "").suffix.lower()
    raw = upload.read()
    try:
        fields = extract_docx_fields(raw) if suffix == ".docx" else []
        return jsonify({
            "filename": upload.filename,
            "kind": suffix.lstrip(".").upper(),
            "fields": fields,
            "message": "已识别 DOCX 占位符" if suffix == ".docx" else "PDF 模板已加入；坐标映射将在桌面编辑器中配置",
        })
    except Exception as exc:
        return jsonify({"error": f"模板解析失败：{exc}"}), 400


@app.route("/api/validate", methods=["POST"])
def validate():
    payload = request.get_json(force=True) or {}
    rows = apply_rules(payload.get("rows", []), payload.get("computedFields", []))
    return jsonify({**validate_rows(rows, payload.get("requiredFields", [])), "rows": rows})


@app.route("/api/generate", methods=["POST"])
def generate():
    payload = request.get_json(force=True) or {}
    rows = apply_rules(payload.get("rows", []), payload.get("computedFields", []))
    required = payload.get("requiredFields", [])
    validation = validate_rows(rows, required)
    settings = payload.get("settings", {})
    filename_pattern = settings.get("filenamePattern") or "{{客户简称}}-报价单-{{报价编号}}"
    folder_pattern = settings.get("folderPattern") or "{{客户简称}}/{{报价编号}}"
    templates = payload.get("templates") or ["quote", "attachment"]

    archive_buffer = io.BytesIO()
    generated = []
    with zipfile.ZipFile(archive_buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for index in validation["validIndexes"]:
            row = rows[index]
            folder = "/".join(safe_component(part) for part in render_pattern(folder_pattern, row).split("/") if part)
            base = safe_component(render_pattern(filename_pattern, row))
            record_files = []
            if "quote" in templates:
                pdf = io.BytesIO()
                make_quote_pdf(row, pdf, settings.get("signature", ""))
                path = f"{folder}/{base}.pdf"
                archive.writestr(path, pdf.getvalue())
                record_files.append(path)
            if "attachment" in templates:
                pdf = io.BytesIO()
                make_attachment_pdf(row, pdf)
                path = f"{folder}/{base}-项目附件.pdf"
                archive.writestr(path, pdf.getvalue())
                record_files.append(path)
            generated.append({"record": row.get("客户简称") or row.get("客户名称"), "files": record_files})

        report = io.StringIO()
        writer = csv.writer(report)
        writer.writerow(["数据行", "记录", "缺失字段"])
        for issue in validation["issues"]:
            writer.writerow([issue["row"], issue["record"], "、".join(issue["missing"])])
        archive.writestr("校验报告.csv", "\ufeff" + report.getvalue())
        manifest = {
            "product": "DocFlow Local",
            "generatedAt": datetime.now().isoformat(timespec="seconds"),
            "privacy": "All processing completed locally.",
            "summary": {"records": len(generated), "files": sum(len(item["files"]) for item in generated), "skipped": validation["invalid"]},
            "items": generated,
        }
        archive.writestr("交付清单.json", json.dumps(manifest, ensure_ascii=False, indent=2))

    archive_buffer.seek(0)
    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    download_name = f"DocFlow-交付包-{stamp}.zip"
    try:
        response = send_file(
            archive_buffer,
            mimetype="application/zip",
            as_attachment=True,
            download_name=download_name,
        )
    except TypeError:  # Flask 1.x
        archive_buffer.seek(0)
        response = send_file(
            archive_buffer,
            mimetype="application/zip",
            as_attachment=True,
            attachment_filename=download_name,
        )
    response.headers["X-DocFlow-Generated"] = str(len(generated))
    response.headers["X-DocFlow-Skipped"] = str(validation["invalid"])
    return response


if __name__ == "__main__":
    print("\n  DocFlow Local 已启动：http://127.0.0.1:4173")
    print("  所有文件只在本机内存中处理。按 Ctrl+C 停止。\n")
    app.run(host="127.0.0.1", port=4173, debug=False)
